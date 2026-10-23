"""Synthetic-only reviewed import contract and real HTTP integration tests."""

from datetime import date, time
from io import BytesIO
import json
import socket
import threading
import time as clock

import httpx
from openpyxl import Workbook
import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker
import uvicorn

from app.core.config import Settings, get_settings
from app.core.database import get_db_session
from app.main import create_app
from app.models.asset_snapshot import AssetSnapshot
from app.models.transaction import Transaction
from app.models.upload_log import UploadLog
from app.models.user_profile_snapshot import UserProfileSnapshot
from app.schemas.upload_review import (
    ReviewApplyRequest,
    ReviewContext,
    ReviewDecision,
    ReviewRevision,
)
from app.services import upload_review_service as svc
from app.services.upload_review_planner import build_review_entries

DAY = date(2026, 10, 23)
CONTEXT = {
    "owner_profile": "finance",
    "account_hash": "a" * 64,
    "message_id": "1abc",
    "archive_hash": "b" * 64,
}


@pytest.fixture(autouse=True)
def no_real_settings(monkeypatch):
    monkeypatch.setitem(Settings.model_config, "env_file", None)
    monkeypatch.setenv("API_KEY", "test-api-key")
    monkeypatch.setenv("DATABASE_URL", "sqlite+aiosqlite:///:memory:")
    get_settings.cache_clear()


def row(name="synthetic", amount=-100, seconds=0, **kw):
    return {
        "date": DAY,
        "time": time(12, 0, seconds),
        "type": "지출",
        "category_major": "기타",
        "category_minor": None,
        "description": name,
        "merchant": name,
        "amount": amount,
        "currency": "KRW",
        "payment_method": None,
        "memo": None,
        **kw,
    }


def tx(**kw):
    return Transaction(**row(**kw))


def workbook(rows=None, asset=100, full_snapshots=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "가계부 내역"
    ws.append(
        [
            "날짜",
            "시간",
            "타입",
            "대분류",
            "소분류",
            "내용",
            "금액",
            "화폐",
            "결제수단",
            "메모",
        ]
    )
    for r in rows if rows is not None else [row()]:
        ws.append(
            [
                r[k]
                for k in (
                    "date",
                    "time",
                    "type",
                    "category_major",
                    "category_minor",
                    "description",
                    "amount",
                    "currency",
                    "payment_method",
                    "memo",
                )
            ]
        )
    snap = wb.create_sheet("뱅샐현황")
    if full_snapshots:
        snap.append([None, "1.고객정보"])
        snap.append([None, None, "synthetic", 20, 700])
        snap.append([None, "2.현금흐름현황"])
    for i in range(20):
        values = [None] * 10
        if i in (0, 8, 12, 16):
            values[1] = {
                0: "3.재무현황",
                8: "4.보험현황",
                12: "5.투자현황",
                16: "6.대출현황",
            }[i]
        if i == 4 and asset is not None:
            values[1], values[2], values[4] = "현금", "synthetic-bank", asset
        # ensure ten columns are physically represented
        values[9] = "" if i != 19 else "padding"
        if full_snapshots and i == 11:
            values = [
                None,
                "synthetic-insurer",
                "insurance",
                None,
                "active",
                100,
                DAY,
                DAY,
                None,
                None,
            ]
        if full_snapshots and i == 15:
            values = [
                None,
                "stock",
                "synthetic-broker",
                "investment",
                None,
                100,
                110,
                10,
                None,
                None,
            ]
        if full_snapshots and i == 19:
            values = [
                None,
                "loan",
                "synthetic-lender",
                "loan",
                None,
                1000,
                900,
                3,
                DAY,
                date(2030, 1, 1),
            ]
        snap.append(values)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


async def prepare(db, data=None):
    return await svc.prepare_review(
        db,
        file_bytes=data or workbook(),
        filename="synthetic.xlsx",
        snapshot_date=DAY,
        context=ReviewContext(**CONTEXT),
    )


def apply_request(plan, **overrides):
    return ReviewApplyRequest(
        **{
            "confirmation": True,
            "plan_digest": plan.plan_digest,
            "context": CONTEXT,
            "decisions": [d.model_dump() for d in plan.decisions],
            "snapshot_effect_ack": True,
            "approval_reference": "hermes://finance/session/synthetic/message/42",
            **overrides,
        }
    )


async def apply(db, plan, tmp_path, data=None, request=None, **kw):
    return await svc.apply_review(
        db,
        plan.operation_id,
        file_bytes=data or workbook(),
        filename="synthetic.xlsx",
        snapshot_date=kw.pop("snapshot_date", DAY),
        request=request or apply_request(plan),
        upload_dir=tmp_path / "uploads",
        **kw,
    )


async def test_atomic_apply_canonical_readback_exact_replay(db_session, tmp_path):
    data = workbook()
    p = await prepare(db_session, data)
    assert (await prepare(db_session, data)).model_dump() == p.model_dump()
    assert p.counts.new == 1 and len(p.snapshot_effects["tables"]) == 5
    done = await apply(db_session, p, tmp_path, data)
    assert done.status == "applied" and done.result.snapshot_verified
    assert done.result.snapshot_counts["asset_snapshots"] == 1
    assert (
        await svc.get_review(db_session, p.operation_id)
    ).model_dump() == done.model_dump()
    assert (
        await apply(db_session, p, tmp_path, data)
    ).model_dump() == done.model_dump()
    assert await db_session.scalar(select(func.count()).select_from(Transaction)) == 1
    assert await db_session.scalar(select(func.count()).select_from(UploadLog)) == 1
    assert list((tmp_path / "uploads").iterdir())
    with pytest.raises(svc.ReviewError, match="different request"):
        await apply(
            db_session,
            p,
            tmp_path,
            data,
            request=apply_request(p, approval_reference="hermes://finance/other/43"),
        )


@pytest.mark.parametrize("change", ["file", "date", "context", "digest", "decisions"])
async def test_exact_binding_rejects_without_financial_write(
    db_session, tmp_path, change
):
    data = workbook()
    p = await prepare(db_session, data)
    request = apply_request(p)
    kw = {}
    if change == "file":
        data += b"different"
    if change == "date":
        kw["snapshot_date"] = date(2026, 10, 24)
    if change == "context":
        request.context.account_hash = "c" * 64
    if change == "digest":
        request.plan_digest = "0" * 64
    if change == "decisions":
        request.decisions[0].action = "keep"
    with pytest.raises(svc.ReviewError):
        await apply(db_session, p, tmp_path, data, request=request, **kw)
    assert await db_session.scalar(select(func.count()).select_from(Transaction)) == 0
    assert await db_session.scalar(select(func.count()).select_from(UploadLog)) == 0


async def test_revise_invalidates_old_digest_and_prepare_preserves_revision(
    db_session, tmp_path
):
    data = workbook()
    p = await prepare(db_session, data)
    revised = await svc.revise_review(
        db_session,
        p.operation_id,
        ReviewRevision(
            plan_digest=p.plan_digest,
            decisions=[ReviewDecision(entry_id="1", action="keep", target_id=None)],
        ),
    )
    assert revised.version == 2 and revised.plan_digest != p.plan_digest
    assert (await prepare(db_session, data)).model_dump() == revised.model_dump()
    with pytest.raises(svc.ReviewError):
        await apply(db_session, p, tmp_path, data)
    done = await apply(db_session, revised, tmp_path, data)
    assert done.result.kept == ["1"] and not done.result.applied
    assert await db_session.scalar(select(func.count()).select_from(Transaction)) == 0


@pytest.mark.parametrize("mutation", ["source", "metadata", "snapshot"])
async def test_full_state_staleness_is_durable(db_session, tmp_path, mutation):
    existing = tx()
    db_session.add(existing)
    await db_session.commit()
    data = workbook()
    p = await prepare(db_session, data)
    if mutation == "source":
        existing.description = "changed"
    if mutation == "metadata":
        existing.memo = "human-private-note"
    if mutation == "snapshot":
        db_session.add(
            AssetSnapshot(
                snapshot_date=DAY,
                side="asset",
                category="현금",
                product_name="other",
                amount=999,
            )
        )
    await db_session.commit()
    with pytest.raises(svc.ReviewError, match="changed"):
        await apply(db_session, p, tmp_path, data)
    assert (await svc.get_review(db_session, p.operation_id)).status == "stale"
    assert await db_session.scalar(select(func.count()).select_from(UploadLog)) == 0


async def test_all_eight_classes_and_occurrences(db_session):
    existing = [
        tx(name="exact", amount=-1),
        tx(name="category", amount=-2),
        tx(name="clock", amount=-3),
        tx(name="old-name", amount=-4),
        tx(name="dup", amount=-5),
        tx(name="amb", amount=-6, seconds=1),
        tx(name="amb", amount=-6, seconds=2),
        tx(name="missing", amount=-7),
    ]
    db_session.add_all(existing)
    await db_session.commit()
    rows = [
        row("exact", -1),
        row("category", -2, category_major="changed"),
        row("clock", -3, seconds=3),
        row("new-name", -4),
        row("dup", -5),
        row("dup", -5),
        row("amb", -6),
        row("new", -99),
    ]
    entries = await build_review_entries(db_session, rows)
    assert {e.change_type.value for e in entries} == set(
        svc.UploadPreviewChangeCounts.model_fields
    )
    assert [e.entry_id for e in entries] == [str(i + 1) for i in range(len(entries))]
    assert entries[4].change_type == entries[5].change_type == "possible_duplicate"
    assert entries[4].candidates == entries[5].candidates
    assert entries[4].recommended_action == entries[5].recommended_action == "keep"
    assert entries[6].recommended_action == "keep"
    assert entries[-1].allowed_actions == ["keep"]


async def test_overlap_never_becomes_false_new_and_null_sentinel_distinct(db_session):
    existing = tx(name="same")
    db_session.add(existing)
    await db_session.commit()
    entries = await build_review_entries(
        db_session, [row("same", category_major="A"), row("same", category_major="B")]
    )
    assert [e.change_type.value for e in entries] == ["source_fields_changed"] * 2
    assert all(
        e.candidates[0].transaction_id == existing.id and e.recommended_action == "keep"
        for e in entries
    )
    entries = await build_review_entries(
        db_session, [row("same", category_minor="<null>")]
    )
    assert entries[0].change_type == "source_fields_changed"
    assert entries[0].field_changes[0].existing_value is None


async def test_update_retains_user_fields_and_missing_row_exactly_unchanged(
    db_session, tmp_path
):
    existing, missing = tx(name="update", amount=-1), tx(name="missing", amount=-2)
    existing.memo, existing.merchant, existing.category_major_user = (
        "user note",
        "user merchant",
        "override",
    )
    db_session.add_all([existing, missing])
    await db_session.commit()
    before = dict(
        (
            await db_session.execute(
                select(Transaction.__table__).where(Transaction.id == missing.id)
            )
        )
        .mappings()
        .one()
    )
    data = workbook([row("update", -1, category_major="corrected")])
    p = await prepare(db_session, data)
    done = await apply(db_session, p, tmp_path, data)
    await db_session.refresh(existing)
    assert done.result.applied[0].transaction_id == existing.id
    assert (
        existing.memo,
        existing.merchant,
        existing.category_major_user,
        existing.category_major,
    ) == ("user note", "user merchant", "override", "corrected")
    after = dict(
        (
            await db_session.execute(
                select(Transaction.__table__).where(Transaction.id == missing.id)
            )
        )
        .mappings()
        .one()
    )
    assert before == after


async def test_duplicate_update_targets_and_partial_invalid_batch_rejected(
    db_session, tmp_path
):
    existing = tx(name="same")
    db_session.add(existing)
    await db_session.commit()
    data = workbook([row("same", category_major="A"), row("same", category_major="B")])
    p = await prepare(db_session, data)
    for decisions in (
        [ReviewDecision(entry_id="1", action="update", target_id=existing.id)],
        [
            ReviewDecision(entry_id=str(i), action="update", target_id=existing.id)
            for i in (1, 2)
        ],
        [
            ReviewDecision(entry_id="1", action="add", target_id=None),
            ReviewDecision(entry_id="2", action="update", target_id=99999),
        ],
    ):
        with pytest.raises(svc.ReviewError):
            await svc.revise_review(
                db_session,
                p.operation_id,
                ReviewRevision(plan_digest=p.plan_digest, decisions=decisions),
            )
    assert (await svc.get_review(db_session, p.operation_id)).version == 1
    assert await db_session.scalar(select(func.count()).select_from(UploadLog)) == 0
    assert await db_session.scalar(select(func.count()).select_from(Transaction)) == 1


async def test_explicit_distinct_duplicate_adds_preserve_occurrences(
    db_session, tmp_path
):
    data = workbook([row(), row()])
    p = await prepare(db_session, data)
    assert p.counts.possible_duplicate == 2
    revised = await svc.revise_review(
        db_session,
        p.operation_id,
        ReviewRevision(
            plan_digest=p.plan_digest,
            decisions=[
                ReviewDecision(entry_id=str(i), action="add", target_id=None)
                for i in (1, 2)
            ],
        ),
    )
    done = await apply(db_session, revised, tmp_path, data)
    assert len({a.transaction_id for a in done.result.applied}) == 2


async def test_snapshots_older_gate_and_missing_profile_preserved(db_session, tmp_path):
    profile = UserProfileSnapshot(
        snapshot_date=DAY, gender="synthetic", age=20, credit_score_kcb=700
    )
    db_session.add(profile)
    await db_session.commit()
    data = workbook()
    p = await prepare(db_session, data)
    assert p.snapshot_effects["same_day_replacement"]
    done = await apply(db_session, p, tmp_path, data)
    assert done.result.snapshot_counts["user_profile_snapshots"] == 1
    await db_session.refresh(profile)
    assert profile.credit_score_kcb == 700
    db_session.add(
        AssetSnapshot(
            snapshot_date=date(2026, 10, 24),
            side="asset",
            category="현금",
            product_name="later",
            amount=1,
        )
    )
    await db_session.commit()
    with pytest.raises(svc.ReviewError, match="older"):
        await prepare(db_session, data)


async def test_snapshot_failure_rolls_back_transactions_and_audit(
    db_session, tmp_path, monkeypatch
):
    data = workbook()
    p = await prepare(db_session, data)

    async def fail(*args, **kw):
        raise RuntimeError("synthetic injected failure")

    monkeypatch.setattr(svc, "verify_snapshots", fail)
    with pytest.raises(RuntimeError):
        await apply(db_session, p, tmp_path, data)
    assert await db_session.scalar(select(func.count()).select_from(Transaction)) == 0
    assert await db_session.scalar(select(func.count()).select_from(UploadLog)) == 0
    assert await db_session.scalar(select(func.count()).select_from(AssetSnapshot)) == 0
    assert (await svc.get_review(db_session, p.operation_id)).status == "prepared"


async def test_http_auth_validation_prepare_apply_readback(
    async_client, api_headers, tmp_path, monkeypatch
):
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "retained"))
    data = workbook()
    path = "/api/v1/upload/review"
    body = {"snapshot_date": DAY.isoformat(), "context": json.dumps(CONTEXT)}
    files = {"file": ("synthetic.xlsx", data)}
    assert (
        await async_client.post(path + "/prepare", data=body, files=files)
    ).status_code == 401
    assert (await async_client.get(path + "/missing")).status_code == 401
    assert (
        await async_client.get(path + "/missing", headers=api_headers)
    ).status_code == 404
    resp = await async_client.post(
        path + "/prepare", data=body, files=files, headers=api_headers
    )
    assert resp.status_code == 200, resp.text
    p = resp.json()
    request = {
        "plan_digest": p["plan_digest"],
        "context": CONTEXT,
        "decisions": p["decisions"],
        "snapshot_effect_ack": True,
        "approval_reference": "hermes://finance/test/42",
    }
    apply_path = path + "/" + p["operation_id"] + "/apply"
    for confirmation in (None, False, 1):
        invalid = {**request}
        if confirmation is not None:
            invalid["confirmation"] = confirmation
        resp = await async_client.post(
            apply_path,
            data={
                "snapshot_date": DAY.isoformat(),
                "apply_request": json.dumps(invalid),
            },
            files=files,
            headers=api_headers,
        )
        assert resp.status_code == 422
    for suffix, kwargs in (
        (
            "/apply",
            {
                "data": {
                    "snapshot_date": DAY.isoformat(),
                    "apply_request": json.dumps({**request, "confirmation": True}),
                },
                "files": files,
            },
        ),
        (
            "/decisions",
            {"json": {"plan_digest": p["plan_digest"], "decisions": p["decisions"]}},
        ),
    ):
        unauthorized = await async_client.post(
            path + "/" + p["operation_id"] + suffix, **kwargs
        )
        assert unauthorized.status_code == 401
    request["confirmation"] = True
    resp = await async_client.post(
        apply_path,
        data={"snapshot_date": DAY.isoformat(), "apply_request": json.dumps(request)},
        files=files,
        headers=api_headers,
    )
    assert resp.status_code == 200, resp.text
    canonical = await async_client.get(
        path + "/" + p["operation_id"], headers=api_headers
    )
    assert canonical.json() == resp.json()


async def test_real_loopback_http_synthetic_import(db_session, tmp_path, monkeypatch):
    """Real TCP HTTP, multipart parser, real workbook and isolated SQLite DB."""
    monkeypatch.setenv("UPLOAD_DIR", str(tmp_path / "real-http-retained"))
    get_settings.cache_clear()
    app = create_app()
    factory = async_sessionmaker(db_session.bind, expire_on_commit=False)

    async def session():
        async with factory() as db:
            yield db

    app.dependency_overrides[get_db_session] = session
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level="error", lifespan="off"))
    thread = threading.Thread(
        target=server.run, kwargs={"sockets": [sock]}, daemon=True
    )
    thread.start()
    try:
        for _ in range(100):
            if server.started:
                break
            clock.sleep(0.02)
        assert server.started
        async with httpx.AsyncClient(
            base_url=f"http://127.0.0.1:{port}", headers={"X-API-Key": "test-api-key"}
        ) as client:
            data = workbook()
            resp = await client.post(
                "/api/v1/upload/review/prepare",
                files={"file": ("synthetic.xlsx", data)},
                data={"snapshot_date": DAY.isoformat(), "context": json.dumps(CONTEXT)},
            )
            assert resp.status_code == 200, resp.text
            p = resp.json()
            request = {
                "confirmation": True,
                "plan_digest": p["plan_digest"],
                "decisions": p["decisions"],
                "context": CONTEXT,
                "snapshot_effect_ack": True,
                "approval_reference": "hermes://finance/synthetic-real-http/42",
            }
            resp = await client.post(
                f"/api/v1/upload/review/{p['operation_id']}/apply",
                files={"file": ("synthetic.xlsx", data)},
                data={
                    "snapshot_date": DAY.isoformat(),
                    "apply_request": json.dumps(request),
                },
            )
            assert resp.status_code == 200, resp.text
            assert resp.json()["result"]["snapshot_verified"]
            readback = await client.get(f"/api/v1/upload/review/{p['operation_id']}")
            assert readback.json() == resp.json()
    finally:
        server.should_exit = True
        thread.join(timeout=10)
        sock.close()
    assert not thread.is_alive()


async def test_all_eight_default_dispositions_apply_without_relabeling(
    db_session, tmp_path
):
    existing = [
        tx(name="exact", amount=-1),
        tx(name="category", amount=-2),
        tx(name="clock", amount=-3),
        tx(name="old-name", amount=-4),
        tx(name="dup", amount=-5),
        tx(name="amb", amount=-6, seconds=1),
        tx(name="amb", amount=-6, seconds=2),
        tx(name="missing", amount=-7),
    ]
    db_session.add_all(existing)
    await db_session.commit()
    rows = [
        row("exact", -1),
        row("category", -2, category_major="changed"),
        row("clock", -3, seconds=3),
        row("new-name", -4),
        row("dup", -5),
        row("dup", -5),
        row("amb", -6),
        row("new", -99),
    ]
    data = workbook(rows)
    p = await prepare(db_session, data)
    assert all(
        getattr(p.counts, key) > 0 for key in svc.UploadPreviewChangeCounts.model_fields
    )
    done = await apply(db_session, p, tmp_path, data)
    assert [(a.entry_id, a.action) for a in done.result.applied] == [
        ("2", "update"),
        ("3", "update"),
        ("8", "add"),
    ]
    assert done.result.kept == ["1", "4", "5", "6", "7", "9"]
    assert [(e.entry_id, e.change_type) for e in done.entries] == [
        (e.entry_id, e.change_type) for e in p.entries
    ]
    assert await db_session.scalar(select(func.count()).select_from(Transaction)) == 9


async def test_populated_five_snapshot_tables_semantic_readback(db_session, tmp_path):
    data = workbook(full_snapshots=True)
    p = await prepare(db_session, data)
    assert set(p.snapshot_effects["incoming_counts"].values()) == {1}
    done = await apply(db_session, p, tmp_path, data)
    assert done.result.snapshot_counts == {name: 1 for name in svc.SNAPSHOT_MODELS}
    assert done.result.snapshot_verified
    assert (
        await db_session.scalar(
            select(func.count()).select_from(
                svc.Base.metadata.tables["asset_source_runs"]
            )
        )
        >= 1
    )


@pytest.mark.parametrize(
    "kind", ["possible_replacement", "possible_duplicate", "ambiguous"]
)
@pytest.mark.parametrize("action", ["keep", "add", "update"])
async def test_explicit_review_class_dispositions(db_session, tmp_path, kind, action):
    if kind == "possible_replacement":
        existing, incoming = [tx(name="old")], row("different")
    elif kind == "possible_duplicate":
        existing, incoming = [tx(), tx()], row()
    else:
        existing, incoming = [tx(seconds=1), tx(seconds=2)], row()
    db_session.add_all(existing)
    await db_session.commit()
    data = workbook([incoming])
    p = await prepare(db_session, data)
    assert p.entries[0].change_type == kind
    assert p.decisions[0].action == "keep"
    selected = ReviewDecision(
        entry_id="1",
        action=action,
        target_id=existing[0].id if action == "update" else None,
    )
    revised = await svc.revise_review(
        db_session,
        p.operation_id,
        ReviewRevision(plan_digest=p.plan_digest, decisions=[selected]),
    )
    done = await apply(db_session, revised, tmp_path, data)
    assert done.result.kept == (["1"] if action == "keep" else [])
    assert await db_session.scalar(
        select(func.count()).select_from(Transaction)
    ) == len(existing) + (action == "add")
    for original in existing:
        await db_session.refresh(original)
        assert (
            original.source_lifecycle_status != "superseded"
            and original.superseded_by_transaction_id is None
        )
    if action == "update":
        assert done.result.applied[0].transaction_id == existing[0].id


@pytest.mark.parametrize("inactive", ["deleted", "merged", "superseded"])
async def test_inactive_update_target_refused(db_session, inactive):
    existing = tx()
    db_session.add(existing)
    await db_session.flush()
    if inactive == "deleted":
        existing.is_deleted = True
    elif inactive == "merged":
        existing.merged_into_id = existing.id
    else:
        existing.source_lifecycle_status = "superseded"
    await db_session.commit()
    p = await prepare(db_session, workbook([row(category_major="changed")]))
    assert p.decisions[0].action == "keep"
    assert "update" not in p.entries[0].allowed_actions
    with pytest.raises(svc.ReviewError):
        await svc.revise_review(
            db_session,
            p.operation_id,
            ReviewRevision(
                plan_digest=p.plan_digest,
                decisions=[
                    ReviewDecision(entry_id="1", action="update", target_id=existing.id)
                ],
            ),
        )


async def test_retention_failure_is_warning_and_exact_replay_not_reimport(
    db_session, tmp_path, monkeypatch
):
    data = workbook()
    p = await prepare(db_session, data)
    calls = []

    def fail(**kw):
        calls.append(kw["upload_id"])
        raise OSError("synthetic private filename must not be logged")

    monkeypatch.setattr(svc, "persist_original_upload", fail)
    done = await apply(db_session, p, tmp_path, data)
    assert done.status == "applied"
    log = await db_session.get(UploadLog, done.result.upload_id)
    assert "retention failed" in log.error_message
    assert (
        await apply(db_session, p, tmp_path, data)
    ).model_dump() == done.model_dump()
    assert calls == [done.result.upload_id]


async def test_update_preserves_links_and_classification(db_session, tmp_path):
    from app.models.loan_account import LoanAccount
    from app.models.loan_transaction_link import LoanTransactionLink
    from app.models.installment_plan import InstallmentPlan
    from app.models.installment_transaction_link import InstallmentTransactionLink
    from app.models.purchase_gate_review import PurchaseGateReview
    from app.models.settlement_group import SettlementMatch

    existing, refund = tx(), tx(name="refund", amount=100)
    existing.cost_kind = "fixed"
    existing.fixed_cost_necessity = "essential"
    existing.spend_necessity = "essential"
    existing.cost_classification_source = "manual"
    existing.recurring_payment_kind = "subscription"
    account = LoanAccount(lender="synthetic", product_name="loan")
    installment = InstallmentPlan(
        display_name="synthetic",
        merchant="synthetic",
        total_installments=2,
        monthly_amount=100,
        first_payment_date=DAY,
    )
    db_session.add_all([existing, refund, account, installment])
    await db_session.flush()
    links = [
        LoanTransactionLink(
            transaction_id=existing.id, loan_account_id=account.id, memo="keep"
        ),
        InstallmentTransactionLink(
            transaction_id=existing.id,
            installment_plan_id=installment.id,
            installment_number=1,
            memo="keep",
        ),
        PurchaseGateReview(
            candidate_key=f"transaction:{existing.id}",
            candidate_type="synthetic",
            transaction_id=existing.id,
            review_status="pending",
            memo="keep",
        ),
        SettlementMatch(
            original_transaction_id=existing.id,
            settlement_transaction_id=refund.id,
            status="user_confirmed",
            matched_amount=100,
        ),
    ]
    db_session.add_all(links)
    await db_session.commit()
    before = [
        (
            model.__table__,
            [
                dict(r)
                for r in (await db_session.execute(select(model.__table__))).mappings()
            ],
        )
        for model in (
            LoanTransactionLink,
            InstallmentTransactionLink,
            PurchaseGateReview,
            SettlementMatch,
        )
    ]
    data = workbook([row(category_major="changed")])
    p = await prepare(db_session, data)
    await apply(db_session, p, tmp_path, data)
    for table, rows in before:
        assert [
            dict(r) for r in (await db_session.execute(select(table))).mappings()
        ] == rows
    await db_session.refresh(existing)
    assert (
        existing.cost_kind,
        existing.fixed_cost_necessity,
        existing.spend_necessity,
        existing.cost_classification_source,
        existing.recurring_payment_kind,
    ) == ("fixed", "essential", "essential", "manual", "subscription")
